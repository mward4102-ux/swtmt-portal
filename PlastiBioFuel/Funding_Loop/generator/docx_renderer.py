# -*- coding: utf-8 -*-
"""DOCX renderer for PlastiBioFuel funding documents (python-docx), matched to the PDF design."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import tokens as T

DISP = "Aptos Display"
BODY = "Aptos"
PAGE_W = Inches(8.5)
MARGIN = Inches(1.0)
CW = Inches(6.5)


def _rgb(h): return RGBColor.from_string(h)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":")), v)
    return e


def shade(el_pr, hexcolor):
    el_pr.append(_el("w:shd", **{"w_val": "clear", "w_color": "auto", "w_fill": hexcolor}))


def cell_bg(cell, hexcolor):
    shade(cell._tc.get_or_add_tcPr(), hexcolor)


def cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for name, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{name}")
        if spec:
            color, sz = spec
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)


def cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for name, v in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:w"), str(int(v * 20))); e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def run(p, text, size=10.5, bold=False, color=T.INK, font=BODY, spacing=None,
        shading=None, italic=False, caps=False):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = _rgb(color)
    rPr = r._element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    rPr.insert(0, rf)
    if spacing:
        rPr.append(_el("w:spacing", w_val=str(int(spacing * 20))))
    if shading:
        shade(rPr, shading)
    if caps:
        rPr.append(_el("w:caps", w_val="1"))
    return r


def para(container, space_before=0, space_after=0, line=1.18, indent=0, keep=False):
    p = container.add_paragraph() if not hasattr(container, "paragraphs") or isinstance(container, Document().__class__) else container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent: pf.left_indent = Pt(indent)
    if keep: pf.keep_with_next = True
    return p


def cell_para(cell, first=True, **kw):
    p = cell.paragraphs[0] if (first and cell.paragraphs and not cell.paragraphs[0].runs) else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(kw.get("space_before", 0)); pf.space_after = Pt(kw.get("space_after", 0))
    pf.line_spacing = kw.get("line", 1.18)
    if kw.get("align") == "right": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if kw.get("align") == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def bare_table(doc, rows, cols, widths, indent_twips=0):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblPr.append(_el("w:tblInd", w_w=str(indent_twips), w_type="dxa"))
    tblPr.append(_el("w:tblLayout", w_type="fixed"))
    for r in t.rows:
        for i, c in enumerate(r.cells):
            c.width = widths[i]
            cell_borders(c)
            cell_margins(c)
    return t


RICH = re.compile(r'<font[^>]*name="Helvetica-Bold"[^>]*>(.*?)</font>|<i>(.*?)</i>|<font[^>]*>(.*?)</font>')


def rich(p, text, size=10.5, color=T.INK, font=BODY):
    """Render the small HTML subset used by the shared content model."""
    pos = 0
    for m in RICH.finditer(text):
        if m.start() > pos:
            run(p, _plain_txt(text[pos:m.start()]), size=size, color=color, font=font)
        if m.group(1) is not None:
            run(p, _plain_txt(m.group(1)), size=size, bold=True, color=color, font=font)
        elif m.group(2) is not None:
            run(p, _plain_txt(m.group(2)), size=size, italic=True, color=color, font=font)
        else:
            col = re.search(r'color="#([0-9A-Fa-f]{6})"', m.group(0))
            bold = 'Helvetica-Bold' in m.group(0)
            run(p, _plain_txt(m.group(3)), size=size, bold=bold,
                color=(col.group(1) if col else color), font=font)
        pos = m.end()
    if pos < len(text):
        run(p, _plain_txt(text[pos:]), size=size, color=color, font=font)


def _plain_txt(s):
    return (s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
             .replace("<br/>", "\n"))


def build(blocks, path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55); sec.bottom_margin = Inches(0.7)
    sec.left_margin = MARGIN; sec.right_margin = MARGIN
    n = doc.styles["Normal"]
    n.font.name = BODY; n.font.size = Pt(10.5); n.font.color.rgb = _rgb(T.INK)
    n.paragraph_format.line_spacing = 1.18; n.paragraph_format.space_after = Pt(6)
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), BODY)
    n.element.get_or_add_rPr().insert(0, rf)

    for b in blocks:
        k = b[0]

        if k == "masthead":
            t = bare_table(doc, 1, 2, [Inches(4.25), Inches(4.25)], indent_twips=-1440)
            for i, (txt, col, sz, sp, al) in enumerate([
                    ("PLASTIBIOFUEL", "FFFFFF", 11, 1.8, "left"),
                    (b[1], T.BLUE_LT, 8, 1.9, "right")]):
                c = t.cell(0, i)
                cell_bg(c, T.BLUE)
                cell_margins(c, top=7, bottom=7, left=(72 if i == 0 else 0), right=(0 if i == 0 else 72))
                p = cell_para(c, align=al, line=1.0)
                run(p, txt, size=sz, bold=True, color=col, font=DISP, spacing=sp)
            r = bare_table(doc, 1, 1, [Inches(8.5)], indent_twips=-1440)
            c = r.cell(0, 0); cell_bg(c, T.PINK); cell_margins(c)
            pp = cell_para(c, line=1.0); pp.paragraph_format.space_after = Pt(0)
            run(pp, "", size=2)
            para(doc, space_after=12, line=1.0)

        elif k == "hero":
            p = para(doc, space_after=5, line=1.1, keep=True)
            run(p, b[1], size=9, bold=True, color=T.BLUE, font=DISP, spacing=1.4)
            for i, line in enumerate(b[2].split("\n")):
                q = para(doc, space_after=(6 if i == len(b[2].split("\n")) - 1 else 0), line=1.03, keep=True)
                run(q, line, size=28, bold=True, color=T.INK, font=DISP)
            if len(b) > 3 and b[3]:
                s = para(doc, space_after=13, line=1.3)
                run(s, b[3], size=11, color=T.SOFT)

        elif k in ("stats", "statbar"):
            items = b[1]; n_ = len(items); w = Inches(6.5 / n_)
            t = bare_table(doc, 1, n_, [w] * n_)
            for i, it in enumerate(items):
                c = t.cell(0, i)
                cell_borders(c, top=(T.INK, 12), bottom=(T.RULE, 6),
                             right=((T.RULE, 6) if i < n_ - 1 else None))
                cell_margins(c, top=8, bottom=8, left=(0 if i == 0 else 12), right=12)
                if k == "stats":
                    label, val, note = it
                    p1 = cell_para(c, line=1.0); run(p1, label, size=7.5, bold=True, color=T.MUTED, font=DISP, spacing=1.1)
                    p2 = cell_para(c, first=False, space_before=4, line=1.0)
                    run(p2, val, size=23, bold=True, color=T.BLUE, font=DISP)
                    p3 = cell_para(c, first=False, space_before=4, line=1.1)
                    run(p3, note, size=9, color=T.MUTED)
                else:
                    val, label = it
                    longest = max(len(v) for v, _ in items)
                    vsize = 17 if longest <= 5 else (13 if longest <= 20 else 11)
                    p1 = cell_para(c, line=1.1); run(p1, val, size=vsize, bold=True, color=T.BLUE, font=DISP)
                    p2 = cell_para(c, first=False, space_before=4, line=1.1)
                    run(p2, label, size=7.5, bold=True, color=T.MUTED, font=DISP, spacing=1.1)
            para(doc, space_after=9, line=1.0)

        elif k == "legend":
            items = b[1]; n_ = len(items); w = Inches(6.5 / n_)
            t = bare_table(doc, 1, n_, [w] * n_)
            for i, (kind, txt) in enumerate(items):
                label, fg, bg = T.PILLS[kind]
                c = t.cell(0, i)
                cell_borders(c, top=(T.INK, 12), bottom=(T.RULE, 6))
                cell_margins(c, top=8, bottom=8, right=12)
                p = cell_para(c, line=1.2)
                run(p, f" {label} ", size=7, bold=True, color=fg, font=DISP, spacing=1.0, shading=bg)
                run(p, "   " + txt, size=8.5, color=T.MUTED)
            para(doc, space_after=9, line=1.0)

        elif k == "alert":
            t = bare_table(doc, 1, 1, [CW])
            c = t.cell(0, 0); cell_bg(c, T.PINK_BG)
            cell_borders(c, left=(T.PINK, 24))
            cell_margins(c, top=7, bottom=7, left=14, right=14)
            p = cell_para(c, line=1.25); run(p, b[1], size=10.5)
            para(doc, space_after=9, line=1.0)

        elif k == "factgrid":
            items = b[1]; rows = (len(items) + 1) // 2
            t = bare_table(doc, rows, 2, [Inches(3.25), Inches(3.25)])
            for i, (label, txt) in enumerate(items):
                c = t.cell(i // 2, i % 2)
                cell_borders(c, top=(T.RULE, 6))
                cell_margins(c, top=7, bottom=7, right=(20 if i % 2 == 0 else 0))
                p1 = cell_para(c, line=1.1); run(p1, label, size=7.5, bold=True, color=T.BLUE, font=DISP, spacing=1.1)
                p2 = cell_para(c, first=False, space_before=3, line=1.24); run(p2, txt, size=9.5)
            para(doc, space_after=4, line=1.0)

        elif k == "factwide":
            t = bare_table(doc, 1, 1, [CW])
            c = t.cell(0, 0); cell_borders(c, top=(T.RULE, 6)); cell_margins(c, top=7, bottom=7)
            p1 = cell_para(c, line=1.1); run(p1, b[1], size=7.5, bold=True, color=T.BLUE, font=DISP, spacing=1.1)
            p2 = cell_para(c, first=False, space_before=3, line=1.26); rich(p2, b[2], size=9.5)
            para(doc, space_after=6, line=1.0)

        elif k == "chip":
            t = bare_table(doc, 1, 2, [Inches(3.6), Inches(2.9)])
            c = t.cell(0, 0); cell_bg(c, T.BLUE); cell_margins(c, top=5, bottom=5, left=10, right=10)
            p = cell_para(c, line=1.1); run(p, b[1], size=9.5, bold=True, color="FFFFFF", font=DISP, spacing=1.2)
            c2 = t.cell(0, 1); cell_margins(c2, top=6, bottom=5, left=12)
            p2 = cell_para(c2, align="right", line=1.1)
            if len(b) > 2 and b[2]:
                run(p2, b[2], size=9, bold=True, color=T.MUTED, font=DISP)
            para(doc, space_after=8, line=1.0)

        elif k == "qblock":
            num, qt, pill, paras, addon = b[1], b[2], b[3], b[4], (b[5] if len(b) > 5 else None)
            t = bare_table(doc, 1, 3, [Inches(0.86), Inches(4.5), Inches(1.14)])
            c0 = t.cell(0, 0); cell_margins(c0, top=2)
            p0 = cell_para(c0, line=0.95); run(p0, num, size=30, bold=True, color=T.GHOST, font=DISP)
            c1 = t.cell(0, 1); cell_margins(c1, top=8, bottom=6); cell_borders(c1, bottom=(T.RULE, 6))
            p1 = cell_para(c1, line=1.15); run(p1, qt, size=14, bold=True, color=T.BLUE, font=DISP)
            c2 = t.cell(0, 2); cell_margins(c2, top=9, bottom=6); cell_borders(c2, bottom=(T.RULE, 6))
            p2 = cell_para(c2, align="right", line=1.2)
            run(p2, f" {pill} ", size=7, bold=True, color="FFFFFF", font=DISP, spacing=1.0, shading=T.BLUE)
            para(doc, space_after=3, line=1.0)
            for lead, txt in paras:
                p = para(doc, space_after=7, line=1.3, indent=62)
                if lead:
                    run(p, lead + " ", size=10.5, bold=True, color=T.BLUE, font=DISP)
                run(p, txt, size=10.5)
            if addon:
                at = bare_table(doc, 1, 2, [Inches(0.86), Inches(5.64)])
                ac = at.cell(0, 1); cell_borders(ac, left=(T.RULE, 12)); cell_margins(ac, left=12, top=3, bottom=3)
                ap = cell_para(ac, line=1.1); run(ap, addon[0], size=7, bold=True, color=T.MUTED, font=DISP, spacing=1.0)
                ap2 = cell_para(ac, first=False, space_before=3, line=1.3); run(ap2, addon[1], size=9.5, color=T.SOFT)
                para(doc, space_after=4, line=1.0)
            para(doc, space_after=8, line=1.0)

        elif k == "step":
            t = bare_table(doc, 1, 2, [Inches(0.45), Inches(6.05)])
            c0 = t.cell(0, 0)
            cell_borders(c0, top=(T.BLUE, 8), bottom=(T.BLUE, 8), left=(T.BLUE, 8), right=(T.BLUE, 8))
            cell_margins(c0, top=4, bottom=4)
            p0 = cell_para(c0, align="center", line=1.0)
            run(p0, str(b[1]), size=10.5, bold=True, color=T.BLUE, font=DISP)
            c1 = t.cell(0, 1); cell_margins(c1, left=14, top=3)
            p1 = cell_para(c1, line=1.3); rich(p1, b[2], size=10.5)
            para(doc, space_after=7, line=1.0)

        elif k == "card":
            kind, ttl, where, who, note = b[1], b[2], b[3], b[4], b[5]
            label, fg, bg = T.PILLS[kind]
            t = bare_table(doc, 1, 2, [Inches(1.55), Inches(4.95)])
            c0 = t.cell(0, 0); cell_borders(c0, bottom=(T.RULE_LT, 4)); cell_margins(c0, top=9, bottom=9, right=12)
            p0 = cell_para(c0, line=1.2)
            run(p0, f" {label} ", size=7, bold=True, color=fg, font=DISP, spacing=1.0, shading=bg)
            c1 = t.cell(0, 1); cell_borders(c1, bottom=(T.RULE_LT, 4)); cell_margins(c1, top=8, bottom=9)
            p1 = cell_para(c1, line=1.2); run(p1, ttl, size=10.5, bold=True, font=DISP)
            p2 = cell_para(c1, first=False, space_before=3, line=1.2)
            run(p2, "WHERE ", size=7, bold=True, color=T.FAINT, font=DISP, spacing=1.0)
            run(p2, where + "   |   ", size=8.5, color=T.MUTED)
            run(p2, "SIGNS ", size=7, bold=True, color=T.FAINT, font=DISP, spacing=1.0)
            run(p2, who, size=8.5, color=T.MUTED)
            p3 = cell_para(c1, first=False, space_before=3, line=1.28)
            run(p3, note, size=9.5, color=T.SOFT)

        elif k == "callout":
            t = bare_table(doc, 1, 1, [CW])
            c = t.cell(0, 0); cell_bg(c, T.PINK_BG); cell_borders(c, left=(T.PINK, 24))
            cell_margins(c, top=10, bottom=10, left=16, right=16)
            p1 = cell_para(c, line=1.15); run(p1, b[1], size=8.5, bold=True, color=T.PINK, font=DISP, spacing=1.2)
            for i_, ln in enumerate([x for x in b[2].split("\n") if x.strip()]):
                p2 = cell_para(c, first=False, space_before=(5 if i_ == 0 else 4), line=1.3)
                run(p2, ln, size=10.5)
            para(doc, space_after=10, line=1.0)

        elif k == "tint":
            t = bare_table(doc, 1, 1, [CW])
            c = t.cell(0, 0); cell_bg(c, T.TINT); cell_margins(c, top=11, bottom=11, left=18, right=18)
            p1 = cell_para(c, line=1.15); run(p1, b[1], size=8.5, bold=True, color=T.BLUE, font=DISP, spacing=1.2)
            if len(b) > 2 and b[2]:
                p2 = cell_para(c, first=False, space_before=4, line=1.25); run(p2, b[2], size=8.5, color=T.MUTED)
            if len(b) > 3 and b[3]:
                for i, ln in enumerate(l for l in str(b[3]).split("\n") if l.strip()):
                    p3 = cell_para(c, first=False, space_before=5 if i == 0 else 7, line=1.32)
                    p3.paragraph_format.left_indent = Pt(10)
                    pPr = p3._p.get_or_add_pPr()
                    bd = OxmlElement("w:pBdr")
                    lb = OxmlElement("w:left")
                    lb.set(qn("w:val"), "single"); lb.set(qn("w:sz"), "12")
                    lb.set(qn("w:space"), "6"); lb.set(qn("w:color"), T.BLUE)
                    bd.append(lb); pPr.append(bd)
                    run(p3, ln, size=10)
            para(doc, space_after=10, line=1.0)

        elif k == "kvrows":
            items = b[1]
            t = bare_table(doc, len(items), 2, [Inches(1.85), Inches(4.65)])
            for i, (kk, vv) in enumerate(items):
                c0 = t.cell(i, 0); c1 = t.cell(i, 1)
                bot = (T.RULE, 6) if i == len(items) - 1 else None
                cell_borders(c0, top=(T.RULE, 6), bottom=bot); cell_borders(c1, top=(T.RULE, 6), bottom=bot)
                cell_margins(c0, top=7, bottom=7, right=14); cell_margins(c1, top=7, bottom=7)
                run(cell_para(c0, line=1.2), kk, size=9.5, bold=True, font=DISP)
                run(cell_para(c1, line=1.28), vv, size=9.5)
            para(doc, space_after=6, line=1.0)

        elif k == "ref":
            idx, cite, doi, why = b[1], b[2], b[3], b[4]
            t = bare_table(doc, 1, 2, [Inches(0.47), Inches(6.03)])
            c0 = t.cell(0, 0); cell_borders(c0, bottom=(T.RULE_LT, 4)); cell_margins(c0, top=4, bottom=11)
            run(cell_para(c0, line=1.1), str(idx), size=15, bold=True, color=T.GHOST, font=DISP)
            c1 = t.cell(0, 1); cell_borders(c1, bottom=(T.RULE_LT, 4)); cell_margins(c1, top=4, bottom=11)
            rich(cell_para(c1, line=1.3), cite, size=10)
            if doi:
                run(cell_para(c1, first=False, space_before=4, line=1.2), doi, size=8.5, bold=True,
                    color=T.BLUE, font=DISP)
            pw = cell_para(c1, first=False, space_before=5, line=1.3)
            pw.paragraph_format.left_indent = Pt(10)
            pPr = pw._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); lb = OxmlElement("w:left")
            lb.set(qn("w:val"), "single"); lb.set(qn("w:sz"), "8")
            lb.set(qn("w:space"), "6"); lb.set(qn("w:color"), T.RULE)
            bd.append(lb); pPr.append(bd)
            run(pw, why, size=9.5, color=T.SOFT)
            para(doc, space_after=5, line=1.0)

        elif k == "h1":
            p = para(doc, space_before=12, space_after=4, line=1.2, keep=True)
            run(p, b[1], size=13, bold=True, color=T.BLUE, font=DISP)
        elif k == "h2":
            p = para(doc, space_before=9, space_after=2, line=1.2, keep=True)
            run(p, b[1], size=11, bold=True, color=T.BLUE, font=DISP)
        elif k == "p":
            p = para(doc, space_after=7, line=1.3)
            if len(b) > 2 and b[2]:
                run(p, b[2], size=10.5, bold=True, font=DISP)
            run(p, b[1], size=10.5)
        elif k == "small":
            p = para(doc, space_after=5, line=1.3)
            rich(p, b[1], size=9.5, color=T.SOFT)
        elif k == "footer":
            t = bare_table(doc, 1, 2, [Inches(4.3), Inches(2.2)])
            c0 = t.cell(0, 0); c1 = t.cell(0, 1)
            cell_borders(c0, top=(T.RULE, 6)); cell_borders(c1, top=(T.RULE, 6))
            cell_margins(c0, top=7); cell_margins(c1, top=7)
            run(cell_para(c0, line=1.25), b[1], size=8, color=T.MUTED)
            run(cell_para(c1, align="right", line=1.25), b[2], size=8, color=T.MUTED)

    doc.save(path)
